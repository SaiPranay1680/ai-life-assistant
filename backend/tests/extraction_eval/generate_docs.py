"""Generate synthetic EVAL-001 PDF/DOCX fixtures (OCR-friendly text)."""

from __future__ import annotations

import sys
from pathlib import Path

import pymupdf
from docx import Document

DOCUMENTS_DIR = Path(__file__).resolve().parent / "documents"


def _write_pdf(path: Path, pages: list[str]) -> None:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page(width=595, height=842)  # A4
        # insert_text keeps a real text layer so read_document_pages uses PDF text (not OCR).
        y = 72
        for line in text.splitlines():
            page.insert_text((50, y), line[:110], fontsize=11)
            y += 16
            if y > 800:
                break
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()


def _write_docx(path: Path, pages: list[str], *, as_table: bool = False) -> None:
    document = Document()
    for index, page_text in enumerate(pages):
        if index:
            document.add_page_break()
        lines = [line for line in page_text.splitlines() if line.strip()]
        if as_table and lines:
            # First line as title paragraph; remaining as a 2-column Field|Value table when ":" present.
            document.add_paragraph(lines[0])
            rows = []
            for line in lines[1:]:
                if ":" in line:
                    key, value = line.split(":", 1)
                    rows.append((key.strip(), value.strip()))
                else:
                    document.add_paragraph(line)
            if rows:
                table = document.add_table(rows=len(rows), cols=2)
                for row_index, (key, value) in enumerate(rows):
                    table.rows[row_index].cells[0].text = key
                    table.rows[row_index].cells[1].text = value
        else:
            for line in lines:
                document.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


DOCUMENTS: dict[str, list[str]] = {
    # Matches DOC-005 utility_bill E2E sample.
    "utility_bill_001.pdf": [
        "\n".join(
            [
                "UTILITY BILL",
                "Bengaluru Power & Light Services",
                "Bill Number: BPL-2026-091842",
                "Customer ID: CUST-458921",
                "Billing Period: 01 Aug 2026 - 31 Aug 2026",
                "Due Date: 15 Sep 2026",
                "Amount Due: INR 1,245.00",
                "Service Address: 42 Example Road, Bengaluru, Karnataka",
            ]
        )
    ],
    # Variation: slash dates, missing customer_id, INR spacing.
    "utility_bill_002.pdf": [
        "\n".join(
            [
                "ELECTRICITY BILL",
                "Mysore Metro Power Co",
                "Bill Number: MMP-77821",
                "Account Number: ACC-554433",
                "Billing Period: 01/09/2026 - 30/09/2026",
                "Due Date: 12/10/2026",
                "Amount Due: INR 980.50",
                "Service Address: 9 Lake View, Mysuru",
            ]
        )
    ],
    "insurance_001.pdf": [
        "\n".join(
            [
                "HEALTH INSURANCE POLICY",
                "SYNTHETIC TEST DOCUMENT — FOR EVAL-001",
                "Provider: SecureCare Insurance",
                "Policy Number: SCI-HLTH-2026-778812",
                "Policy Holder: Test Customer",
                "Start Date: 01 Sep 2026",
                "Expiry Date: 31 Aug 2027",
                "Premium: INR 18,500.00",
            ]
        )
    ],
    "insurance_002.pdf": [
        "\n".join(
            [
                "Car Insurance Policy",
                "Motor vehicle cover document",
                "Provider: RoadShield Motors",
                "Policy Number: CAR-9911-XX",
                "Policy Holder: Eval Driver",
                "Premium: INR 12,500.00",
                "Effective Date: 05/03/2026",
                "Expiry Date: 04/03/2027",
            ]
        )
    ],
    "purchase_receipt_001.pdf": [
        "\n".join(
            [
                "PURCHASE RECEIPT",
                "Merchant: TechWorld Electronics",
                "Receipt Number: TW-REC-2026-004581",
                "Purchase Date: 18 Sep 2026",
                "Item 1: 27-inch Monitor — Qty 1 — INR 24,000",
                "Item 2: Mechanical Keyboard — Qty 1 — INR 8,000",
                "Item 3: Wireless Mouse — Qty 2 — INR 5,000 each",
                "Subtotal: INR 42,000.00",
                "Tax: INR 7,560.00",
                "Total: INR 49,560.00",
                "Payment Method: Credit Card",
            ]
        )
    ],
    # OCR rupee-as-I amounts + missing payment_method.
    "purchase_receipt_002.pdf": [
        "\n".join(
            [
                "PURCHASE RECEIPT",
                "Merchant: ElectroMart Retail",
                "Receipt Number: EM-8821",
                "Purchase Date: 2026-09-18",
                "Item 1: USB Hub — Qty 1 — I1,200",
                "Subtotal: I1,200.00",
                "Tax: I216.00",
                "Total: I1,416.00",
            ]
        )
    ],
    "warranty_001.pdf": [
        "\n".join(
            [
                "WARRANTY CERTIFICATE",
                "Product: UltraView 27-inch 4K Monitor",
                "Brand: UltraView",
                "Model: UV-27-4K-2026",
                "Serial Number: UV27SN884291",
                "Warranty Provider: TechWorld Electronics",
                "Purchase Date: 18 Sep 2026",
                "Warranty Start: 18 Sep 2026",
                "Warranty Expiry: 17 Sep 2028",
                "Warranty Duration: 24 months",
            ]
        )
    ],
    # Variation: slash dates, optional brand omitted, multi-page terms.
    "warranty_002.pdf": [
        "\n".join(
            [
                "WARRANTY CARD",
                "Product: NoiseBuds Pro",
                "Model: NB-P-11",
                "Serial Number: NB11-445566",
                "Warranty Provider: SoundWave Retail",
                "Purchase Date: 01/02/2026",
                "Warranty Start: 01/02/2026",
                "Warranty Expiry: 01/02/2027",
                "Warranty Duration: 12 months",
            ]
        ),
        "Terms: Coverage excludes physical damage. Keep this card with your purchase receipt.",
    ],
    "generic_001.pdf": [
        "\n".join(
            [
                "PASSPORT",
                "Republic of Example",
                "Document Date: 10 Jan 2024",
                "Holder: Eval User",
            ]
        )
    ],
    "generic_002.pdf": [
        "\n".join(
            [
                "Personal Note",
                "Meeting notes for apartment lease renewal discussion.",
                "Document Date: 22/08/2026",
                "Provider: Home Office Records",
            ]
        )
    ],
}

# DOCX mirrors of the primary fixtures (same labeled text as PDFs — no DOCX-only rules).
DOCX_DOCUMENTS: dict[str, tuple[list[str], bool]] = {
    "utility_bill_003.docx": (DOCUMENTS["utility_bill_001.pdf"], False),
    "insurance_003.docx": (DOCUMENTS["insurance_001.pdf"], False),
    "purchase_receipt_003.docx": (DOCUMENTS["purchase_receipt_001.pdf"], False),
    "warranty_003.docx": (DOCUMENTS["warranty_001.pdf"], False),
}


def generate_all() -> list[Path]:
    written: list[Path] = []
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    for name, pages in DOCUMENTS.items():
        path = DOCUMENTS_DIR / name
        _write_pdf(path, pages)
        written.append(path)
    for name, (pages, as_table) in DOCX_DOCUMENTS.items():
        path = DOCUMENTS_DIR / name
        _write_docx(path, pages, as_table=as_table)
        written.append(path)
    return written


def main() -> int:
    paths = generate_all()
    print(f"Wrote {len(paths)} documents to {DOCUMENTS_DIR}")
    for path in paths:
        print(f"  - {path.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
