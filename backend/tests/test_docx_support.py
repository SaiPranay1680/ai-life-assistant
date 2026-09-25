"""DOC-003: DOCX validation, extraction, and pipeline integration tests."""

from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from app.modules.documents.docx_extract import (
    DOCX_MIME,
    DocxExtractionError,
    extract_docx_pages,
    extract_docx_text,
    validate_docx_file,
)
from app.modules.documents.extract import build_structured_payload, guess_fields
from app.modules.documents.validation import (
    detect_mime,
    validate_extension,
    validate_extraction,
    validate_file_type,
)


def _write_docx(path: Path, paragraphs: list[str] | None = None, table_rows: list[tuple[str, str]] | None = None) -> None:
    document = Document()
    for line in paragraphs or []:
        document.add_paragraph(line)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=2)
        for index, (left, right) in enumerate(table_rows):
            table.rows[index].cells[0].text = left
            table.rows[index].cells[1].text = right
    document.save(path)


class DocxValidationTests(unittest.TestCase):
    def test_extension_and_mime_allowed(self) -> None:
        self.assertEqual(validate_extension("bill.docx"), ".docx")
        validate_file_type(".docx", DOCX_MIME)

    def test_detect_mime_from_docx_bytes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.docx"
            _write_docx(path, ["Hello"])
            data = path.read_bytes()
            self.assertEqual(detect_mime(data), DOCX_MIME)

    def test_valid_docx_container(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "ok.docx"
            _write_docx(path, ["Line one"])
            validate_docx_file(path)

    def test_empty_docx_extracts_safely(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "empty.docx"
            Document().save(path)
            pages = extract_docx_pages(path)
            self.assertEqual(pages, [(1, "")])

    def test_paragraphs_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "paras.docx"
            _write_docx(path, ["First paragraph", "", "Second paragraph"])
            text = extract_docx_text(path)
            self.assertIn("First paragraph", text)
            self.assertIn("Second paragraph", text)

    def test_tables_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "table.docx"
            _write_docx(
                path,
                paragraphs=["UTILITY BILL"],
                table_rows=[("Provider", "City Power"), ("Amount Due", "INR 100.00")],
            )
            text = extract_docx_text(path)
            self.assertIn("UTILITY BILL", text)
            self.assertIn("Provider", text)
            self.assertIn("City Power", text)
            self.assertIn("Amount Due", text)

    def test_malformed_docx_raises(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.docx"
            path.write_bytes(b"PK\x03\x04not-a-real-zip-payload")
            with self.assertRaises(DocxExtractionError):
                validate_docx_file(path)

    def test_zip_without_word_parts_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "plain.zip.docx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("readme.txt", "hello")
            with self.assertRaises(DocxExtractionError):
                validate_docx_file(path)

    def test_executable_renamed_to_docx_rejected_by_mime(self) -> None:
        data = b"MZ\x90\x00not-a-docx"
        self.assertEqual(detect_mime(data), "application/octet-stream")
        with self.assertRaises(ValueError):
            validate_file_type(".docx", detect_mime(data))

    def test_pdf_renamed_to_docx_rejected(self) -> None:
        data = b"%PDF-1.4\n%EOF\n"
        self.assertEqual(detect_mime(data), "application/pdf")
        with self.assertRaises(ValueError):
            validate_file_type(".docx", detect_mime(data))
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "fake.docx"
            path.write_bytes(data)
            with self.assertRaises(DocxExtractionError):
                validate_docx_file(path)

    def test_image_renamed_to_docx_rejected(self) -> None:
        data = b"\xff\xd8\xff\xe0" + b"\x00" * 32
        self.assertEqual(detect_mime(data), "image/jpeg")
        with self.assertRaises(ValueError):
            validate_file_type(".docx", detect_mime(data))

    def test_wrong_mime_extension_combinations(self) -> None:
        with self.assertRaises(ValueError):
            validate_file_type(".pdf", DOCX_MIME)
        with self.assertRaises(ValueError):
            validate_file_type(".docx", "application/pdf")
        with self.assertRaises(ValueError):
            validate_file_type(".docx", "image/png")

    def test_existing_pdf_png_jpeg_still_allowed(self) -> None:
        validate_extension("a.pdf")
        validate_extension("a.jpg")
        validate_extension("a.jpeg")
        validate_extension("a.png")
        validate_file_type(".pdf", "application/pdf")
        validate_file_type(".jpg", "image/jpeg")
        validate_file_type(".png", "image/png")
        self.assertEqual(detect_mime(b"%PDF-1.7"), "application/pdf")
        self.assertEqual(detect_mime(b"\xff\xd8\xff\xdb"), "image/jpeg")
        self.assertEqual(detect_mime(b"\x89PNG\r\n\x1a\n"), "image/png")


class DocxPipelineTests(unittest.TestCase):
    def test_docx_classification_and_doc005_extraction(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bill.docx"
            _write_docx(
                path,
                paragraphs=[
                    "UTILITY BILL",
                    "Bengaluru Power & Light Services",
                    "Bill Number: BPL-2026-091842",
                    "Customer ID: CUST-458921",
                    "Billing Period: 01 Aug 2026 - 31 Aug 2026",
                    "Due Date: 15 Sep 2026",
                    "Amount Due: INR 1,245.00",
                    "Service Address: 42 Example Road, Bengaluru, Karnataka",
                ],
            )
            pages = extract_docx_pages(path)
            fields = guess_fields(pages)
            self.assertEqual(fields["documentType"].raw, "utility_bill")
            self.assertEqual(fields["provider"].raw, "Bengaluru Power & Light Services")
            self.assertEqual(fields["bill_number"].raw, "BPL-2026-091842")
            payload = build_structured_payload("utility_bill", fields)
            validated = validate_extraction("utility_bill", payload)
            self.assertEqual(validated.due_date.normalized, "2026-09-15")
            self.assertEqual(validated.amount_due.normalized, "1245")

    def test_docx_purchase_receipt_through_pipeline(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "receipt.docx"
            _write_docx(
                path,
                paragraphs=[
                    "PURCHASE RECEIPT",
                    "Merchant: TechWorld Electronics",
                    "Receipt Number: TW-REC-2026-004581",
                    "Purchase Date: 18 Sep 2026",
                    "Item 1: 27-inch Monitor — Qty 1 — INR 24,000",
                    "Subtotal: INR 42,000.00",
                    "Tax: INR 7,560.00",
                    "Total: INR 49,560.00",
                    "Payment Method: Credit Card",
                ],
            )
            pages = extract_docx_pages(path)
            fields = guess_fields(pages)
            self.assertEqual(fields["documentType"].raw, "purchase_receipt")
            self.assertFalse(bool(fields.get("transaction_date") and fields["transaction_date"].raw))
            validated = validate_extraction(
                "purchase_receipt",
                build_structured_payload("purchase_receipt", fields),
            )
            self.assertEqual(validated.merchant.raw, "TechWorld Electronics")
            self.assertEqual(validated.total.normalized, "49560")
            self.assertIsNone(validated.transaction_date)


if __name__ == "__main__":
    unittest.main()
